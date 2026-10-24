from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.7)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].font.bold = True
for name, size in [('Heading 1', 15), ('Heading 2', 12)]:
    s = doc.styles[name]
    s.font.size = Pt(size)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(13)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True

def clean(text):
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    return text.replace('`', '')

def text_runs(p, text):
    for i, piece in enumerate(clean(text).split('**')):
        r = p.add_run(piece)
        r.bold = bool(i % 2)

def table(rows):
    t = doc.add_table(rows=1, cols=len(rows[0]))
    t.style = 'Table Grid'
    t.autofit = False
    # Keep each numerical table readable with wrapping headings and stable widths.
    n = len(rows[0])
    widths = ([1.65] + [(7.1-1.65)/(n-1)]*(n-1)) if n > 2 else [1.45, 5.65]
    for c,w in zip(t.columns,widths):
        c.width = Inches(w)
    for j, row in enumerate(rows):
        cells = t.rows[0].cells if j == 0 else t.add_row().cells
        for c, value, w in zip(cells,row,widths):
            c.width = Inches(w)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.line_spacing = 1.08
            text_runs(p,value)
            for r in p.runs:
                r.font.size = Pt(10 if n > 5 else 10.5)
                if j == 0: r.bold = True
            if j == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'EDF1F4'); c._tc.get_or_add_tcPr().append(shade)
        trpr = t.rows[j]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if j == 0: trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

source = (ROOT/'RESULTS-SUMMARY.md').read_text(encoding='utf-8')
source = source.split('## 8.')[0]
lines = source.splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-:]+', x) for x in cells): rows.append(cells)
            i += 1
        table(rows); continue
    if line.startswith('!['):
        target = re.search(r'\]\(([^)]+)\)',line).group(1)
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ROOT/target),width=Inches(7.05))
        cap = doc.add_paragraph(style='Caption')
        cap.add_run('图1 分层后的水流与40点最大速度对照。左图黑色为障碍，点颜色为速度误差；右图黑点为实测、橙叉为基线、蓝点为分层结果。底图为计算水深，不能视为实测淹没范围。')
    elif line.startswith('# '):
        doc.add_paragraph('HYRA洪水工作台',style='Title')
        doc.add_paragraph('Merewether对照实验结果报告',style='Subtitle')
    elif line.startswith('## '):
        title = re.sub(r'^\d+\.\s*','',line[3:]).replace('／','与')
        doc.add_paragraph(title,style='Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:],style='Heading 2')
    else:
        p = doc.add_paragraph()
        text_runs(p,line)
    i += 1
doc.add_paragraph('资料来源与核验说明',style='Heading 1')
for line in [
    '数据来源为Ayog与Sharifian于2022年发布的配套数据包 Test 4，许可为CC BY 4.0。公开地址 https://zenodo.org/records/6914888 。',
    '参考论文为Kesserwani等于2023年发表于Journal of Hydraulic Engineering的 Shallow-flow velocity predictions using discontinuous Galerkin solutions。DOI 10.1061/JHEND8.HYENG-13244。',
    '边界参考源码为官方LISFLOOD-FP 8.0发布包，DOI 10.5281/zenodo.4073011。已校验发布包MD5，并编译原边界函数做受控探针；尚未获得2023年论文实际运行二进制或精确提交版本。',
    '本报告数值来自五组完整800秒运行的归档数据。各组最大流速RMSE已从全部40点重新计算核对；分层改善比例按未四舍五入数值计算。图1已嵌入本文件，阅读报告无需访问项目目录。',
    '本次输入核验和边界探针没有新增全域洪水计算。Harvey真实水位工作台属于观测接入与回放，尚无对应完整事件的求解精度结果，未纳入本报告评分。',
]:
    doc.add_paragraph(line)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('HYRA  |  ');r.font.size=Pt(9)
f = OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title = 'Merewether对照实验结果报告'
doc.core_properties.subject = '洪水工作台物理实验验证及输入核验'
doc.core_properties.author = 'HYRA项目组'
out = ROOT/'Merewether洪水对照实验结果报告.docx'
doc.save(out)
print(out)
